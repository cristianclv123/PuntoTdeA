"""Validación temporal de personalización de campañas desde archivos Excel."""

from __future__ import annotations

import hashlib
import json
import re
import zipfile
from datetime import date, datetime, timedelta

import openpyxl
from django.db import transaction
from django.utils import timezone
from openpyxl.utils.exceptions import InvalidFileException
from rest_framework.exceptions import ValidationError as APIValidationError

from . import meta_templates
from ..models import (
    AudienceSegment,
    BroadcastRecipient,
    Campaign,
    CampaignPersonalization,
    Contact,
)

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_ROWS = 10_000
VALIDATION_TTL_MINUTES = 30
TEMPLATE_VARIABLE_RE = re.compile(r"\{\{(\d+)\}\}")


def _canonical_fingerprint(data: dict) -> str:
    encoded = json.dumps(
        data,
        sort_keys=True,
        ensure_ascii=False,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _params_fingerprint(params: dict) -> str:
    return _canonical_fingerprint(params)


def campaign_fingerprint(data: dict) -> str:
    if isinstance(data, dict):
        get = data.get
    else:
        get = lambda name, default=None: getattr(data, name, default)
    segment = get("segment")
    return _canonical_fingerprint(
        {
            "segment": getattr(segment, "pk", segment),
            "template_meta_id": get("template_meta_id", ""),
            "template_name": get("template_name", ""),
            "template_language": get("template_language", "es_CO"),
            "template_snapshot": get("template_snapshot") or {},
            "default_params": get("default_params") or {},
        }
    )


def _open_workbook(file_obj):
    filename = getattr(file_obj, "name", "")
    if not filename.lower().endswith(".xlsx"):
        raise APIValidationError({"file": "El archivo debe tener extensión .xlsx."})
    if getattr(file_obj, "size", 0) > MAX_FILE_SIZE:
        raise APIValidationError({"file": "El archivo no puede superar 10 MB."})

    try:
        workbook = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError, zipfile.BadZipFile) as exc:
        raise APIValidationError({"file": "No se pudo leer el archivo Excel .xlsx."}) from exc

    return workbook


def _delete_expired_validations() -> None:
    CampaignPersonalization.objects.filter(
        campaign__isnull=True,
        expires_at__lte=timezone.now(),
    ).delete()


def inspect_columns(file_obj) -> list[dict[str, object]]:
    _delete_expired_validations()
    workbook = _open_workbook(file_obj)
    try:
        sheet = workbook.active
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header or not any(value is not None and str(value).strip() for value in header):
            raise APIValidationError({"file": "El archivo no contiene encabezados."})
        return [
            {
                "index": index,
                "label": str(value).strip() if value is not None and str(value).strip() else f"Columna {index + 1}",
            }
            for index, value in enumerate(header)
        ]
    finally:
        workbook.close()


def _cell_text(value, *, document=False) -> str:
    if value is None:
        return ""
    if document and isinstance(value, (int, float)) and not isinstance(value, bool):
        if float(value).is_integer():
            return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    if document and re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def template_variables(body_text: str) -> list[str]:
    return sorted({str(int(index)) for index in TEMPLATE_VARIABLE_RE.findall(body_text or "")}, key=int)


def _render_preview(body_text: str, params: dict[str, str]) -> str:
    return TEMPLATE_VARIABLE_RE.sub(
        lambda match: params.get(str(int(match.group(1))), match.group(0)),
        body_text or "",
    )


def _validate_mapping(variable_mappings, document_column, column_count, variables):
    if not isinstance(variable_mappings, dict):
        raise APIValidationError({"variable_mappings": "El mapeo debe ser un objeto JSON."})
    if isinstance(document_column, bool) or not isinstance(document_column, int):
        raise APIValidationError({"document_column": "Selecciona una columna de documento válida."})
    if document_column < 0 or document_column >= column_count:
        raise APIValidationError({"document_column": "La columna de documento no existe en el archivo."})

    unexpected = set(variable_mappings) - set(variables)
    if unexpected:
        raise APIValidationError(
            {"variable_mappings": f"Hay mapeos para variables inexistentes: {', '.join(sorted(unexpected))}."}
        )

    normalized = {}
    for variable in variables:
        mapping = variable_mappings.get(variable)
        if not isinstance(mapping, dict):
            raise APIValidationError(
                {"variable_mappings": f"Asigna una columna o un valor fijo a la variable {{{{{variable}}}}}."}
            )
        source = mapping.get("source")
        if source == "column":
            column = mapping.get("column")
            if isinstance(column, bool) or not isinstance(column, int) or not 0 <= column < column_count:
                raise APIValidationError(
                    {"variable_mappings": f"La columna asignada a {{{{{variable}}}}} no existe."}
                )
            normalized[variable] = {"source": "column", "column": column}
        elif source == "fixed":
            normalized[variable] = {"source": "fixed", "value": _cell_text(mapping.get("value"))}
        else:
            raise APIValidationError(
                {"variable_mappings": f"El origen de {{{{{variable}}}}} debe ser columna o valor fijo."}
            )
    return normalized


def validate_personalization(
    *,
    file_obj,
    segment: AudienceSegment,
    campaign_data: dict,
    document_column: int,
    variable_mappings: dict,
) -> tuple[CampaignPersonalization, dict]:
    try:
        snapshot = meta_templates.verify_campaign_template(
            campaign_data.get("template_meta_id"),
            campaign_data.get("template_name"),
            campaign_data.get("template_language"),
            campaign_data.get("template_snapshot"),
        )
    except meta_templates.TemplateVerificationError as exc:
        raise APIValidationError({"template": str(exc)}) from exc

    campaign_data["template_snapshot"] = snapshot
    body_text = snapshot["body_text"]
    if not body_text:
        raise APIValidationError({"template_snapshot": "La plantilla no contiene un cuerpo válido."})

    workbook = _open_workbook(file_obj)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or not any(value is not None and str(value).strip() for value in header):
            raise APIValidationError({"file": "El archivo no contiene encabezados."})

        variables = template_variables(body_text)
        mappings = _validate_mapping(
            variable_mappings,
            document_column,
            len(header),
            variables,
        )
        parsed_rows = []
        for row_number, row in enumerate(rows, start=2):
            if row_number > MAX_ROWS + 1:
                raise APIValidationError({"file": f"El archivo no puede superar {MAX_ROWS} filas de datos."})
            if not row or not any(value is not None and str(value).strip() for value in row):
                continue
            raw_document = row[document_column] if document_column < len(row) else None
            document_number = _cell_text(raw_document, document=True)
            parsed_rows.append((row_number, row, document_number))
    finally:
        workbook.close()

    document_counts = {}
    for _, _, document in parsed_rows:
        if document:
            document_counts[document] = document_counts.get(document, 0) + 1

    documents = set(document_counts)
    existing_contacts = {
        contact.document_number: contact
        for contact in Contact.objects.filter(document_number__in=documents)
    }
    segment_contacts = {
        contact.document_number: contact
        for contact in segment.contacts.filter(document_number__in=documents)
    }

    valid_rows = []
    valid_preview = []
    rejected_rows = []
    for row_number, row, document in parsed_rows:
        reasons = []
        if not document:
            reasons.append("Falta el documento obligatorio.")
        elif document_counts[document] > 1:
            reasons.append("El documento está duplicado en el archivo.")
        elif document not in existing_contacts:
            reasons.append("No existe un contacto con ese documento.")
        elif document not in segment_contacts:
            reasons.append("El contacto no pertenece a la audiencia seleccionada.")

        contact = segment_contacts.get(document)
        params = {}
        if contact:
            if not contact.full_name.strip():
                reasons.append("El contacto no tiene nombre.")
            if not contact.phone.strip():
                reasons.append("El contacto no tiene teléfono.")
            if contact.whatsapp_opt_in == Contact.OptInStatus.BAJA:
                reasons.append("El contacto se dio de baja de WhatsApp.")

        for variable, mapping in mappings.items():
            if mapping["source"] == "column":
                column = mapping["column"]
                value = _cell_text(row[column] if column < len(row) else None)
            else:
                value = mapping["value"]
            if not value:
                reasons.append(f"Falta un valor para la variable {{{{{variable}}}}}.")
            params[variable] = value

        if reasons:
            rejected_rows.append(
                {
                    "row": row_number,
                    "document_number": document,
                    "reason": " ".join(reasons),
                }
            )
            continue

        valid_rows.append({"contact_id": str(contact.pk), "params": params})
        valid_preview.append(
            {
                "row": row_number,
                "document_number": document,
                "contact_name": contact.full_name,
                "message": _render_preview(body_text, params),
            }
        )

    now = timezone.now()
    _delete_expired_validations()
    validation = CampaignPersonalization.objects.create(
        segment=segment,
        template_fingerprint=campaign_fingerprint(campaign_data),
        valid_rows=valid_rows,
        valid_count=len(valid_rows),
        rejected_count=len(rejected_rows),
        expires_at=now + timedelta(minutes=VALIDATION_TTL_MINUTES),
    )
    return validation, {
        "validation_id": str(validation.validation_id),
        "valid_count": len(valid_rows),
        "rejected_count": len(rejected_rows),
        "valid_rows": valid_preview,
        "rejected_rows": rejected_rows,
        "expires_at": validation.expires_at.isoformat(),
    }


@transaction.atomic
def create_personalized_campaign(
    campaign_data: dict,
    validation_id,
    confirm_partial: bool,
) -> Campaign:
    now = timezone.now()
    try:
        validation = CampaignPersonalization.objects.select_for_update().get(
            validation_id=validation_id,
            campaign__isnull=True,
            expires_at__gt=now,
        )
    except CampaignPersonalization.DoesNotExist as exc:
        raise APIValidationError(
            {"personalization_validation_id": "La validación no existe, expiró o ya fue utilizada."}
        ) from exc

    if campaign_fingerprint(campaign_data) != validation.template_fingerprint:
        raise APIValidationError(
            {"personalization_validation_id": "La audiencia o la plantilla cambió después de validar el Excel."}
        )
    if validation.segment_id != getattr(campaign_data.get("segment"), "pk", None):
        raise APIValidationError(
            {"personalization_validation_id": "La validación pertenece a otra audiencia."}
        )
    if validation.rejected_count and not confirm_partial:
        raise APIValidationError(
            {
                "confirm_partial": (
                    f"Hay {validation.rejected_count} filas rechazadas; confirma explícitamente "
                    "que continuarás solo con las filas válidas."
                )
            }
        )
    if not validation.valid_rows:
        raise APIValidationError(
            {"personalization_validation_id": "No hay filas válidas para crear la campaña."}
        )

    campaign = Campaign.objects.create(**campaign_data)
    contact_ids = [row.get("contact_id") for row in validation.valid_rows]
    if len(contact_ids) != len(set(contact_ids)):
        raise APIValidationError(
            {"personalization_validation_id": "La validación contiene contactos duplicados."}
        )

    contacts = {
        str(contact.pk): contact
        for contact in campaign.segment.contacts.filter(pk__in=contact_ids)
    }
    if set(contacts) != set(contact_ids):
        raise APIValidationError(
            {"personalization_validation_id": "Cambió la pertenencia de un contacto a la audiencia."}
        )

    required_variables = template_variables(
        (campaign.template_snapshot or {}).get("body_text", "")
    )
    param_fingerprints = {}
    for row in validation.valid_rows:
        contact = contacts[row["contact_id"]]
        params = row.get("params")
        if not isinstance(params, dict) or any(
            not str(params.get(variable, "")).strip() for variable in required_variables
        ):
            raise APIValidationError(
                {"personalization_validation_id": "La validación contiene variables sin valor."}
            )
        param_fingerprints[row["contact_id"]] = _params_fingerprint(params)
        BroadcastRecipient.objects.update_or_create(
            campaign=campaign,
            contact=contact,
            defaults={"phone_snapshot": contact.phone, "params": params},
        )

    validation.campaign = campaign
    validation.authorized_contact_ids = sorted(contact_ids)
    validation.authorized_param_fingerprints = param_fingerprints
    validation.confirmed_at = now
    validation.valid_rows = []
    validation.save(
        update_fields=[
            "campaign",
            "authorized_contact_ids",
            "authorized_param_fingerprints",
            "confirmed_at",
            "valid_rows",
        ]
    )
    return campaign


def assert_authorized_recipients(campaign: Campaign, recipients) -> None:
    try:
        authorization = campaign.personalization
    except CampaignPersonalization.DoesNotExist:
        return

    if authorization.confirmed_at is None:
        raise PersonalizationAuthorizationError("La personalización no está confirmada.")

    if (
        authorization.segment_id != campaign.segment_id
        or campaign_fingerprint(campaign) != authorization.template_fingerprint
    ):
        raise PersonalizationAuthorizationError(
            "La audiencia o la plantilla cambió después de validar el Excel."
        )

    authorized_ids = set(authorization.authorized_contact_ids)
    actual_ids = {str(recipient.contact_id) for recipient in recipients}
    segment_ids = {
        str(contact_id)
        for contact_id in campaign.segment.contacts.filter(pk__in=authorized_ids).values_list("pk", flat=True)
    }
    if actual_ids != authorized_ids or segment_ids != authorized_ids:
        raise PersonalizationAuthorizationError(
            "Los destinatarios de la campaña ya no coinciden con las filas validadas."
        )

    for recipient in recipients:
        expected = authorization.authorized_param_fingerprints.get(str(recipient.contact_id))
        if not expected or _params_fingerprint(recipient.params) != expected:
            raise PersonalizationAuthorizationError(
                "Los parámetros de un destinatario cambiaron después de validar el Excel."
            )


class PersonalizationAuthorizationError(Exception):
    """La campaña ya no coincide con su autorización de destinatarios."""
